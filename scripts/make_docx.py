"""Build the unified manuscript, index, and cover letter as DOCX files.

The converter intentionally supports the Markdown constructs used by the
paper: headings, paragraphs, inline emphasis/code, lists, code fences,
tables, and relative PNG figures. It has no dependency on a Markdown HTML
renderer, which keeps the Word build deterministic.
"""
from __future__ import annotations

import html
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
PAPERS_DIR = ROOT / "papers"
DOCX_DIR = PAPERS_DIR / "docx"
DOCX_DIR.mkdir(parents=True, exist_ok=True)

DOCUMENTS = (
    ("00_INDEX_en.md", "00_INDEX.docx"),
    ("COVER_LETTER.md", "COVER_LETTER.docx"),
    ("preprint_unified_en.md", "preprint_unified_en.docx"),
)


def set_run_font(run, name: str, size: float, bold: bool | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def setup_document(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(19)
    section.bottom_margin = Mm(19)
    section.left_margin = Mm(19)
    section.right_margin = Mm(19)
    section.start_type = WD_SECTION_START.NEW_PAGE

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(4)

    title_style = doc.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(17)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(8)

    for style_name, size, before, after in (
        ("Heading 1", 14, 10, 4),
        ("Heading 2", 12, 8, 3),
        ("Heading 3", 11, 7, 2),
    ):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    doc.core_properties.title = title
    doc.core_properties.author = "AGI Research Kit Contributors"
    doc.core_properties.subject = "Unified submission draft"

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Page ")
    set_run_font(run, "Times New Roman", 9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


INLINE_RE = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\))"
)


def add_inline(paragraph, text: str, size: float = 10.5) -> None:
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(html.unescape(text[position:match.start()]))
            set_run_font(run, "Times New Roman", size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(html.unescape(token[2:-2]))
            set_run_font(run, "Times New Roman", size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(html.unescape(token[1:-1]))
            set_run_font(run, "Courier New", max(size - 1, 8.5))
        elif token.startswith("*"):
            run = paragraph.add_run(html.unescape(token[1:-1]))
            set_run_font(run, "Times New Roman", size)
            run.italic = True
        else:
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            label, target = link.groups() if link else (token, "")
            rendered = f"{label} ({target})" if target else label
            run = paragraph.add_run(html.unescape(rendered))
            set_run_font(run, "Times New Roman", size)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(html.unescape(text[position:]))
        set_run_font(run, "Times New Roman", size)


def set_cell_margins(cell, top: int = 80, start: int = 90,
                     bottom: int = 80, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start),
                          ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    normalized = [row + [""] * (columns - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=columns)
    table.style = "Table Grid"

    max_lengths = [max(4, min(40, max(len(row[i]) for row in normalized)))
                   for i in range(columns)]
    total_weight = sum(max_lengths)
    usable_dxa = 9720
    widths = [max(900, round(usable_dxa * value / total_weight))
              for value in max_lengths]
    widths[-1] += usable_dxa - sum(widths)
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])

    for row_index, row in enumerate(normalized):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value, size=8.5)
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_image(doc: Document, source: Path, alt_text: str) -> None:
    if not source.exists():
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(f"[Missing figure: {source.name}]")
        set_run_font(run, "Times New Roman", 10, bold=True)
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(source), width=Inches(6.15))
    drawing = run._r.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("descr", alt_text)
    paragraph.paragraph_format.keep_with_next = True


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    raw_rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            raw_rows.append(cells)
        index += 1
    return raw_rows, index


def build_docx(markdown_path: Path, output_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8-sig").splitlines()
    title = next((line[2:].strip() for line in lines if line.startswith("# ")),
                 markdown_path.stem)
    doc = Document()
    setup_document(doc, title)

    paragraph_buffer: list[str] = []
    in_code = False
    code_lines: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        paragraph = doc.add_paragraph()
        add_inline(paragraph, " ".join(part.strip() for part in paragraph_buffer))
        paragraph_buffer.clear()

    index = 0
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if stripped.startswith("```"):
            if in_code:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.2)
                paragraph.paragraph_format.space_after = Pt(6)
                run = paragraph.add_run("\n".join(code_lines))
                set_run_font(run, "Courier New", 8.5)
                code_lines.clear()
                in_code = False
            else:
                flush_paragraph()
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image:
            flush_paragraph()
            alt_text, relative_path = image.groups()
            add_image(doc, (markdown_path.parent / relative_path).resolve(), alt_text)
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            table_rows, index = parse_table(lines, index)
            add_table(doc, table_rows)
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if text in {"References", "Appendix A - Hardware Footprint"}:
                doc.add_page_break()
            if level == 1:
                paragraph = doc.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, text, size=17)
                for run in paragraph.runs:
                    run.bold = True
            else:
                paragraph = doc.add_heading(level=min(level - 1, 3))
                add_inline(paragraph, text, size={2: 14, 3: 12, 4: 11}[level])
            index += 1
            continue

        list_item = re.match(r"^[-*]\s+(.+)$", stripped)
        numbered_item = re.match(r"^\d+\.\s+(.+)$", stripped)
        if list_item or numbered_item:
            flush_paragraph()
            style = "List Bullet" if list_item else "List Number"
            paragraph = doc.add_paragraph(style=style)
            add_inline(paragraph, (list_item or numbered_item).group(1))
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            add_inline(paragraph, stripped.lstrip("> "))
            for run in paragraph.runs:
                run.italic = True
            index += 1
            continue

        if not stripped or stripped == "---":
            flush_paragraph()
        else:
            paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    print("=== Generating unified DOCX artifacts ===")
    for source_name, output_name in DOCUMENTS:
        source = PAPERS_DIR / source_name
        output = DOCX_DIR / output_name
        build_docx(source, output)
        print(f"  {output.relative_to(ROOT)} ({output.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
